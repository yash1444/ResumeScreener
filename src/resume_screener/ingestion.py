"""Read resumes from a folder. One bad file never stops the batch."""
from __future__ import annotations

import hashlib
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

from .models import RawResume

log = logging.getLogger(__name__)

SUPPORTED = {".pdf", ".docx", ".txt"}
# zero-width space / non-joiner / joiner / word joiner / BOM: invisible, but they break heading matching
INVISIBLE = {0x200B: None, 0x200C: None, 0x200D: None, 0x2060: None, 0xFEFF: None}
MIN_TEXT_CHARS = 80  # below this a PDF is almost certainly scanned / empty


@dataclass
class IngestionReport:
    resumes: list[RawResume] = field(default_factory=list)
    failed: list[dict] = field(default_factory=list)
    duplicates: list[dict] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)
    total_files: int = 0


def _read_pdf(path: Path) -> tuple[str, list[str], str]:
    """Returns (text in PDF stream order, hyperlinks, text re-ordered by block position)."""
    import fitz  # PyMuPDF

    text_parts: list[str] = []
    alt_parts: list[str] = []
    links: list[str] = []
    with fitz.open(path) as doc:
        for page in doc:
            text_parts.append(page.get_text())
            blocks = [b for b in page.get_text("blocks") if b[6] == 0]
            blocks.sort(key=lambda b: (round(b[1] / 4), b[0]))  # 4pt row bands, then left to right
            alt_parts.append("\n".join(b[4].strip("\n") for b in blocks))
            links.extend(l["uri"] for l in page.get_links() if l.get("uri"))
    return "\n".join(text_parts), links, "\n".join(alt_parts)


def _read_pdf_fallback(path: Path) -> tuple[str, list[str], str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((p.extract_text() or "") for p in reader.pages), [], ""


def _read_docx(path: Path) -> tuple[str, list[str], str]:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells)
    links = [r.target_ref for r in d.part.rels.values() if "hyperlink" in r.reltype]
    return "\n".join(parts), links, ""


def read_resume(path: Path) -> RawResume:
    """Extract text + hyperlinks from one file. Raises on unreadable input."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            text, links, alt = _read_pdf(path)
        except Exception as exc:  # corrupt stream etc. -> try a second parser
            log.warning("PyMuPDF failed for %s (%s); trying pypdf", path.name, exc)
            text, links, alt = _read_pdf_fallback(path)
        if len(text.strip()) < MIN_TEXT_CHARS:  # PyMuPDF may return little on odd fonts
            try:
                other, _, _ = _read_pdf_fallback(path)
                if len(other.strip()) > len(text.strip()):
                    text, alt = other, ""
            except Exception:
                pass
    elif suffix == ".docx":
        text, links, alt = _read_docx(path)
    else:
        text, links, alt = path.read_text(encoding="utf-8", errors="ignore"), [], ""
    text = re.sub(r"[ \t]+\n", "\n", text.replace("\x00", "").translate(INVISIBLE))
    alt = re.sub(r"[ \t]+\n", "\n", alt.replace("\x00", "").translate(INVISIBLE))
    if len(text.strip()) < MIN_TEXT_CHARS:
        raise ValueError("no extractable text (scanned/image-only or empty file)")
    return RawResume(
        path=str(path), filename=path.name, text=text, links=links, alt_text=alt,
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
    )


def _text_fingerprint(text: str) -> str:
    norm = re.sub(r"\W+", " ", text.lower()).strip()
    return hashlib.sha256(norm.encode()).hexdigest()


def ingest_folder(folder: Path) -> IngestionReport:
    report = IngestionReport()
    if not folder.is_dir():
        raise FileNotFoundError(f"Input folder not found: {folder}")
    files = sorted(p for p in folder.rglob("*") if p.is_file())
    seen_bytes: dict[str, str] = {}
    seen_text: dict[str, str] = {}
    for path in files:
        if path.suffix.lower() not in SUPPORTED:
            report.skipped.append(path.name)
            continue
        report.total_files += 1
        try:
            resume = read_resume(path)
        except Exception as exc:
            log.error("Failed to read %s: %s", path.name, exc)
            report.failed.append({"filename": path.name, "error": f"{type(exc).__name__}: {exc}"})
            continue
        fp = _text_fingerprint(resume.text)
        dup_of = seen_bytes.get(resume.sha256) or seen_text.get(fp)
        if dup_of:
            report.duplicates.append({"filename": path.name, "duplicate_of": dup_of})
            continue
        seen_bytes[resume.sha256] = path.name
        seen_text[fp] = path.name
        report.resumes.append(resume)
    return report
